import os
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_summary_report(df, output_path="data/output/monthly_sales_summary.xlsx"):
    """クレンジング済みDataFrameから、集計ピボットと装飾付きのExcelレポートを生成する"""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    # 1. 集計用ピボットテーブルの作成（行: 拠点, 列: カテゴリー, 値: 売上金額の合計）
    pivot_category = pd.pivot_table(
        df,
        index="拠点",
        columns="カテゴリー",
        values="売上金額",
        aggfunc="sum",
        fill_value=0,
        margins=True,
        margins_name="合計"
    )

    # 2. 担当者別の売上集計
    pivot_staff = df.groupby(["拠点", "担当者"])["売上金額"].sum().reset_index()
    pivot_staff = pivot_staff.sort_values(by=["拠点", "売上金額"], ascending=[True, False])

    # 3. ExcelWriterで複数シート出力
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        # シート1: 統合マスタ
        df.to_excel(writer, sheet_name="統合データマスタ", index=False)
        
        # シート2: 集計サマリー
        pivot_category.to_excel(writer, sheet_name="集計サマリー", startrow=2, startcol=1)
        pivot_staff.to_excel(writer, sheet_name="集計サマリー", startrow=2, startcol=8, index=False)

        # openpyxlのワークブックオブジェクトを取得して装飾
        wb = writer.book
        ws_summary = wb["集計サマリー"]
        ws_master = wb["統合データマスタ"]

        # スタイルの定義
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")  # 濃紺
        header_font = Font(name="Yu Gothic UI", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Yu Gothic UI", size=14, bold=True, color="1F4E79")
        data_font = Font(name="Yu Gothic UI", size=10)
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        # 集計サマリーのタイトル設定
        ws_summary["B2"] = "■ 拠点・カテゴリー別 売上集計 (円)"
        ws_summary["B2"].font = title_font
        ws_summary["I2"] = "■ 担当者別 実績集計 (円)"
        ws_summary["I2"].font = title_font

        # シート全体のフォント・列幅自動調整・通貨フォーマット適用
        for ws in [ws_master, ws_summary]:
            ws.views.sheetView[0].showGridLines = True
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value is not None:
                        # 金額列の書式設定（¥#,##0）
                        if isinstance(cell.value, (int, float)) and cell.value > 100:
                            cell.number_format = '"¥"#,##0'
                        # ヘッダー行の装飾
                        if cell.row == 3 and ws == ws_summary:
                            cell.fill = header_fill
                            cell.font = header_font
                            cell.alignment = Alignment(horizontal="center")
                        elif cell.row == 1 and ws == ws_master:
                            cell.fill = header_fill
                            cell.font = header_font
                            cell.alignment = Alignment(horizontal="center")
                        else:
                            if not (ws == ws_summary and cell.row == 2):
                                cell.font = data_font
                                cell.border = thin_border

                        val_str = str(cell.value)
                        max_len = max(max_len, len(val_str))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    print(f"📊 集計レポートを生成しました: {output_path}")

if __name__ == "__main__":
    from cleaner import load_and_clean_all
    cleaned_data = load_and_clean_all()
    create_summary_report(cleaned_data)